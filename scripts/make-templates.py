"""Generate the free downloadable templates served from public/templates/.

    python scripts/make-templates.py

Needs openpyxl and python-docx. Output is committed, so this only has to be
re-run when a template's content changes. Each file carries a small
"Free template from The Counsel" credit, and nothing in them is firm data.
"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = pathlib.Path(__file__).resolve().parent.parent / "public" / "templates"
OUT.mkdir(parents=True, exist_ok=True)

GOLD = "B38A3E"
INK = "1A1D23"
IVORY = "F8F5EF"
CREDIT = "Free template from The Counsel · thecounsels.org"
thin = Side(style="thin", color="D9D2C4")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


# ---------------------------------------------------------------- Excel helpers
def header_row(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub, cols):
    ws.cell(row=1, column=1, value=text).font = Font(bold=True, size=16, color=INK)
    ws.cell(row=2, column=1, value=sub).font = Font(italic=True, size=9, color="7C7468")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)


def body_grid(ws, first, last, cols):
    for r in range(first, last + 1):
        for c in range(1, cols + 1):
            ws.cell(row=r, column=c).border = BOX


def instructions(wb, lines):
    ws = wb.create_sheet("How to use")
    ws.column_dimensions["A"].width = 100
    ws["A1"] = "How to use this template"
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    for i, line in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True)
    ws.cell(row=len(lines) + 4, column=1, value=CREDIT).font = Font(italic=True, size=9, color=GOLD)


# ---------------------------------------------------------------- 1. timesheet
def timesheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"
    cols = ["Date", "Client", "Matter / file ref", "Work done", "Start", "Finish", "Hours", "Billable?", "Hourly rate", "Amount"]
    title(ws, "Lawyer Timesheet", f"Lawyer: ____________________    Week starting: ____________    {CREDIT}", len(cols))
    header_row(ws, 4, cols, [12, 22, 18, 44, 9, 9, 9, 11, 13, 15])
    first, last = 5, 54
    for r in range(first, last + 1):
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        for c in (5, 6):
            ws.cell(row=r, column=c).number_format = "hh:mm"
        ws.cell(row=r, column=7, value=f'=IF(AND(E{r}<>"",F{r}<>""),ROUND((F{r}-E{r})*24,2),"")')
        ws.cell(row=r, column=7).number_format = "0.00"
        ws.cell(row=r, column=10, value=f'=IF(AND(G{r}<>"",H{r}="Yes",I{r}<>""),G{r}*I{r},"")')
        ws.cell(row=r, column=9).number_format = "#,##0"
        ws.cell(row=r, column=10).number_format = "#,##0"
    body_grid(ws, first, last, len(cols))
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H{first}:H{last}")
    t = last + 2
    for label, formula in [
        ("Total hours", f"=SUM(G{first}:G{last})"),
        ("Billable hours", f'=SUMIF(H{first}:H{last},"Yes",G{first}:G{last})'),
        ("Billable amount", f"=SUM(J{first}:J{last})"),
    ]:
        ws.cell(row=t, column=9, value=label).font = Font(bold=True)
        c = ws.cell(row=t, column=10, value=formula)
        c.font = Font(bold=True)
        c.number_format = "#,##0.00"
        t += 1
    ws.freeze_panes = "A5"
    instructions(wb, [
        "1. Fill in your name and the week at the top.",
        "2. Log each piece of work as you finish it, not at the end of the month. Enter Start and Finish as times (e.g. 09:30).",
        "3. Hours and Amount fill in automatically. Amount only counts rows marked Billable = Yes with an hourly rate.",
        "4. Always name the client and the matter: unassigned time can't be billed.",
        "5. Describe the work the way you'd want it to read on an invoice (e.g. 'Drafted statement of claim', not 'drafting').",
        "Tired of re-typing timesheets into invoices? In The Counsel, time logged against a matter goes straight into the invoice.",
    ])
    wb.save(OUT / "lawyer-timesheet-template.xlsx")


# ---------------------------------------------------------------- 2. court diary
def court_diary():
    wb = Workbook()
    ws = wb.active
    ws.title = "Court Diary"
    cols = ["Date", "Time", "Court", "Division / courtroom", "Suit no.", "Matter (parties)", "Business of the day", "Counsel appearing", "Outcome", "Next date", "Notes"]
    title(ws, "Court Diary & Cause List Tracker", f"Rows due in the next 7 days turn gold. Past dates with no outcome turn red.    {CREDIT}", len(cols))
    header_row(ws, 4, cols, [12, 8, 22, 18, 16, 32, 18, 20, 24, 12, 30])
    first, last = 5, 204
    for r in range(first, last + 1):
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=2).number_format = "hh:mm"
        ws.cell(row=r, column=10).number_format = "dd/mm/yyyy"
    body_grid(ws, first, last, len(cols))
    dv = DataValidation(type="list", formula1='"Mention,Hearing,Trial,Ruling,Judgment,Report,Adoption of addresses,Motion,Other"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G{first}:G{last}")
    rng = f"A{first}:K{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND($A{first}<>\"\",$A{first}>=TODAY(),$A{first}<=TODAY()+7)"], fill=PatternFill("solid", fgColor="F3E3B5")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND($A{first}<>\"\",$A{first}<TODAY(),$I{first}=\"\")"], fill=PatternFill("solid", fgColor="F6D0CB")))
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{last}"
    instructions(wb, [
        "1. Add every court date the moment you get it: from the hearing notice, the court's cause list, or the last sitting.",
        "2. After each sitting, fill in the Outcome and the Next date, then add the next date as its own new row.",
        "3. Rows in the next 7 days are highlighted gold, so the week ahead is obvious at a glance.",
        "4. Any past date with no outcome recorded turns red: chase it before the next one comes round.",
        "5. Use the filter on the header row to see one court, one matter or one counsel at a time.",
        "The Counsel does this for you: every hearing on one shared calendar, with email and WhatsApp reminders the day before.",
    ])
    wb.save(OUT / "court-diary-template.xlsx")


# ---------------------------------------------------------------- 3. fee invoice
def fee_invoice():
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    for col, w in zip("ABCDEF", [14, 46, 10, 14, 16, 4]):
        ws.column_dimensions[col].width = w
    ws["A1"] = "[YOUR FIRM NAME]"
    ws["A1"].font = Font(bold=True, size=18, color=INK)
    ws["A2"] = "[Address] · [Phone] · [Email]"
    ws["A2"].font = Font(size=9, color="7C7468")
    ws["D1"] = "INVOICE"
    ws["D1"].font = Font(bold=True, size=20, color=GOLD)
    meta = [("Invoice no.", "INV-0001"), ("Date", ""), ("Due date", ""), ("Matter ref.", "")]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=4, value=k).font = Font(bold=True)
        ws.cell(row=i, column=5, value=v)
    ws["A5"] = "Bill to"
    ws["A5"].font = Font(bold=True, color=GOLD)
    ws["A6"] = "[Client name]"
    ws["A7"] = "[Client address]"
    ws["A8"] = "Re: [Matter description, e.g. Suit No. ___: X v. Y]"

    r = 10
    ws.cell(row=r, column=1, value="Professional fees").font = Font(bold=True, size=12, color=INK)
    r += 1
    header_row(ws, r, ["Date", "Description of work", "Hours", "Rate", "Amount"])
    fees_first = r + 1
    for i in range(8):
        rr = fees_first + i
        ws.cell(row=rr, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=rr, column=5, value=f'=IF(AND(C{rr}<>"",D{rr}<>""),C{rr}*D{rr},"")')
        for c in (4, 5):
            ws.cell(row=rr, column=c).number_format = "#,##0.00"
    fees_last = fees_first + 7
    body_grid(ws, fees_first, fees_last, 5)

    r = fees_last + 2
    ws.cell(row=r, column=1, value="Disbursements (filing fees, service, travel, copies)").font = Font(bold=True, size=12, color=INK)
    r += 1
    header_row(ws, r, ["Date", "Description", "", "", "Amount"])
    dis_first = r + 1
    for i in range(5):
        ws.cell(row=dis_first + i, column=5).number_format = "#,##0.00"
    dis_last = dis_first + 4
    body_grid(ws, dis_first, dis_last, 5)

    r = dis_last + 2
    rows = [
        ("Professional fees", f"=SUM(E{fees_first}:E{fees_last})"),
        ("VAT rate (set for your country)", 0.075),
        ("VAT on fees", f"=E{r}*E{r + 1}"),
        ("Disbursements", f"=SUM(E{dis_first}:E{dis_last})"),
        ("TOTAL DUE", f"=E{r}+E{r + 2}+E{r + 3}"),
        ("Less: retainer / payments received", 0),
        ("BALANCE DUE", f"=E{r + 4}-E{r + 5}"),
    ]
    for i, (label, val) in enumerate(rows):
        lc = ws.cell(row=r + i, column=4, value=label)
        vc = ws.cell(row=r + i, column=5, value=val)
        lc.alignment = Alignment(horizontal="right")
        vc.number_format = "0.0%" if i == 1 else "#,##0.00"
        if label in ("TOTAL DUE", "BALANCE DUE"):
            lc.font = vc.font = Font(bold=True, color=INK, size=12)
            vc.fill = PatternFill("solid", fgColor="F3E3B5")
    r += len(rows) + 1
    for line in [
        "Payment details: [Bank name] · [Account name] · [Account number]",
        "Payment terms: due within [14] days of the invoice date. Please quote the invoice number with your payment.",
        "Thank you for your instructions.",
    ]:
        ws.cell(row=r, column=1, value=line)
        r += 1
    ws.cell(row=r + 1, column=1, value=CREDIT).font = Font(italic=True, size=8, color=GOLD)
    instructions(wb, [
        "1. Replace everything in [square brackets] with your firm's and the client's details.",
        "2. List professional fees line by line: date, what was done, hours and rate. The amount calculates itself.",
        "3. List disbursements separately: clients expect to see filing fees and travel apart from your fees.",
        "4. Set the VAT rate for your country (it's pre-filled at 7.5%, Nigeria's standard rate; change or zero it as needed).",
        "5. Deduct any retainer already received so the balance due is exactly what the client owes.",
        "6. Number invoices in sequence and never reuse a number.",
        "In The Counsel, unbilled time and disbursements for a matter turn into a numbered draft invoice in one click.",
    ])
    wb.save(OUT / "legal-fee-invoice-template.xlsx")


# ---------------------------------------------------------------- Word helpers
def new_doc(title_text, subtitle):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    h = d.add_paragraph()
    run = h.add_run(title_text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(INK)
    s = d.add_paragraph()
    sr = s.add_run(subtitle)
    sr.italic = True
    sr.font.size = Pt(9)
    sr.font.color.rgb = RGBColor.from_string("7C7468")
    return d


def heading(d, text):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(GOLD)
    p.paragraph_format.space_before = Pt(12)


def para(d, text):
    d.add_paragraph(text)


def field(d, label, lines=1):
    p = d.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run("_" * 60)
    for _ in range(lines - 1):
        d.add_paragraph("_" * 86)


def credit(d):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(CREDIT)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GOLD)


# ---------------------------------------------------------------- 4. retainer agreement
def retainer():
    d = new_doc(
        "Letter of Engagement / Retainer Agreement",
        "A starting point for qualified legal practitioners. Adapt every clause to the matter and to the Rules of Professional Conduct "
        "and other rules that apply where you practise. Text in [square brackets] must be completed or removed.",
    )
    para(d, "This agreement is made on [date] between [Firm name], of [address] (\"the Firm\"), and [Client name], of [address] (\"the Client\").")
    sections = [
        ("1. The matter", [
            "The Client instructs the Firm to act in: [describe the matter, e.g. 'the recovery of the debt owed by X Ltd', or 'Suit No. ___ between ___ and ___'].",
        ]),
        ("2. Scope of work", [
            "The Firm will: [list the work, e.g. review documents, advise, draft and file processes, appear in court up to judgment at first instance].",
            "Not included unless agreed in writing: [e.g. appeals, enforcement of judgment, related or new proceedings]. Any further work will be agreed separately.",
        ]),
        ("3. Who will handle the matter", [
            "[Name], [position], will have overall responsibility, assisted by [names/roles]. The Firm may involve other lawyers and staff where appropriate.",
        ]),
        ("4. Professional fees", [
            "Fees will be charged on the following basis: [choose one: a fixed fee of ___ for the scope above / hourly rates of ___ (partner), ___ (associate) / a fee of ___ payable in stages: ___].",
            "Any estimate given is not a fixed quotation unless stated. The Firm will tell the Client promptly if it expects the estimate to be exceeded.",
        ]),
        ("5. Retainer deposit", [
            "The Client will pay a retainer deposit of [amount] before work starts. The Firm will apply it to [fees / disbursements / both] as they fall due and will [top it up by request / account for any balance at the end of the matter].",
        ]),
        ("6. Disbursements", [
            "The Client will reimburse expenses incurred on the matter, such as filing fees, service of processes, certified true copies, expert fees and travel. "
            "Any single disbursement above [amount] will be approved by the Client in advance.",
        ]),
        ("7. Invoices and payment", [
            "The Firm will invoice [monthly / at each stage / on completion]. Invoices are payable within [14] days. "
            "Where invoices remain unpaid, the Firm may, on reasonable notice and subject to its professional obligations, stop work on the matter.",
        ]),
        ("8. The Client's responsibilities", [
            "The Client will give clear and timely instructions, provide complete and accurate information and documents, and tell the Firm promptly of any change in circumstances.",
        ]),
        ("9. Communication", [
            "The Firm will keep the Client informed of significant developments and every court date. The Client's preferred contact is [name, phone, email].",
        ]),
        ("10. Confidentiality and data", [
            "The Firm will keep the Client's affairs confidential, subject to its legal and professional obligations, and will process the Client's personal data in line with the data protection law that applies.",
        ]),
        ("11. Conflicts of interest", [
            "The Firm has checked for conflicts of interest and is not aware of any. If one arises, the Firm will tell the Client and act in line with its professional rules.",
        ]),
        ("12. Ending the engagement", [
            "The Client may end this engagement at any time by written notice. The Firm may stop acting for good reason and on reasonable notice, in line with its professional obligations. "
            "Fees and disbursements incurred up to the end of the engagement remain payable.",
        ]),
        ("13. Files", [
            "The Firm will keep the matter file for at least [number] years after the matter ends, after which it may be destroyed unless the Client asks for it.",
        ]),
        ("14. Governing law", [
            "This agreement is governed by the laws of [country / state].",
        ]),
    ]
    for h, paras in sections:
        heading(d, h)
        for t in paras:
            para(d, t)
    heading(d, "Signed")
    para(d, "For the Firm: ______________________   Name: ______________________   Date: __________")
    para(d, "The Client: ______________________   Name: ______________________   Date: __________")
    credit(d)
    d.save(OUT / "retainer-agreement-template.docx")


# ---------------------------------------------------------------- 5. client intake form
def intake():
    d = new_doc(
        "New Client Intake Form",
        "Complete at the first meeting, before the firm accepts instructions. Run the conflict check in section 3 before any advice is given.",
    )
    heading(d, "1. Client details")
    para(d, "Client type:   ☐ Individual     ☐ Company / organisation")
    for f in ["Full name / company name", "RC number (companies)", "Contact person and position", "Address", "Phone", "Email", "Means of identification seen (type and number)", "Preferred way to be contacted"]:
        field(d, f)
    heading(d, "2. The matter")
    for f, n in [("Type of matter (e.g. litigation, property, corporate, family)", 1), ("What the client wants to achieve", 2), ("Summary of the facts", 4)]:
        field(d, f, n)
    heading(d, "3. Conflict check")
    for f in ["Opposing party / parties", "Their lawyers (if known)", "Other people or companies involved"]:
        field(d, f)
    para(d, "Conflict check completed by: ______________   Date: __________   Result:  ☐ Clear   ☐ Conflict (do not proceed)")
    heading(d, "4. Deadlines and court dates")
    for f in ["Next court date or filing deadline", "Suit number and court (if already in court)", "Any limitation period concerns"]:
        field(d, f)
    heading(d, "5. Documents received")
    para(d, "☐ Contract / agreement     ☐ Correspondence     ☐ Court processes     ☐ Title documents     ☐ Other: ______________")
    heading(d, "6. Fees discussed")
    para(d, "Basis:  ☐ Fixed fee     ☐ Hourly     ☐ Staged     ☐ Retainer deposit")
    for f in ["Amount / rates quoted", "Engagement letter sent on"]:
        field(d, f)
    heading(d, "7. Firm use")
    for f in ["Matter / file reference", "Responsible lawyer", "How the client heard about us"]:
        field(d, f)
    heading(d, "8. Consent")
    para(d, "I confirm the information above is accurate, and I consent to the firm holding and using my personal data to act for me on this matter, in line with the applicable data protection law.")
    para(d, "Signature: ______________________   Name: ______________________   Date: __________")
    credit(d)
    d.save(OUT / "client-intake-form-template.docx")


if __name__ == "__main__":
    timesheet()
    court_diary()
    fee_invoice()
    retainer()
    intake()
    for f in sorted(OUT.iterdir()):
        print(f"{f.name:40s} {f.stat().st_size // 1024} KB")
